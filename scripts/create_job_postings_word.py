import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_callout_box(doc, title, text, bg_hex="FBF9F6", border_hex="0F766E"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    # Set left border thick, clear others
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    r_title = p.add_run(title)
    r_title.bold = True
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(11)
    r_title.font.color.rgb = RGBColor(15, 118, 110)

    for line in text.split("\n\n"):
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(2)
        p2.paragraph_format.space_after = Pt(4)
        p2.paragraph_format.line_spacing = 1.15
        r = p2.add_run(line.strip())
        r.font.name = 'Calibri'
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(50, 45, 40)

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(6)

def create_document():
    doc = docx.Document()

    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(30, 25, 20)

    # Document Header Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("Rivyn — Internshala Job Postings")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(15, 118, 110)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("Talent Acquisition & Internship Postings for Bachelor's Students (India)")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(110, 100, 90)

    # Intro notice
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(12)
    p_intro.paragraph_format.line_spacing = 1.15
    r_intro = p_intro.add_run(
        "Here are the updated job postings with that crucial note prominently highlighted for both roles.\n\n"
        "This will immediately put students at ease, reduce imposter syndrome, and effectively filter out "
        "low-effort ChatGPT spam so you only get genuine, passionate applicants."
    )
    r_intro.font.size = Pt(10.5)
    r_intro.font.italic = True
    r_intro.font.color.rgb = RGBColor(60, 55, 50)

    # ==================== POSTING 1 ====================
    h1_p1 = doc.add_paragraph()
    h1_p1.paragraph_format.space_before = Pt(12)
    h1_p1.paragraph_format.space_after = Pt(4)
    r_h1 = h1_p1.add_run("Posting 1: UI/UX & Frontend Developer Intern (Bachelor’s Students)")
    r_h1.font.name = 'Calibri'
    r_h1.font.size = Pt(16)
    r_h1.bold = True
    r_h1.font.color.rgb = RGBColor(15, 118, 110)

    # Meta Table
    tbl_p1 = doc.add_table(rows=6, cols=2)
    tbl_p1.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_p1.autofit = False
    
    meta_p1 = [
        ("Job / Internship Title", "UI/UX & Frontend Developer Intern (B.Tech / BCA / B.Des)"),
        ("Profile Category", "Web Development / Front End Development / UI/UX Design"),
        ("Workplace Type", "Work from Home (Remote)"),
        ("Duration", "3 to 6 Months (Flexible around college exams & classes)"),
        ("Stipend", "₹15,000 – ₹25,000 / month (+ Pre-Placement Offer upon graduation)"),
        ("Eligible Degrees & Batches", "B.Tech / B.E. / BCA / B.Sc (CS/IT) / B.Des (2025, 2026, 2027, 2028 passouts)"),
    ]

    for idx, (label, val) in enumerate(meta_p1):
        row = tbl_p1.rows[idx]
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(4.3)
        
        set_cell_background(row.cells[0], "F2EEE8")
        set_cell_background(row.cells[1], "FAFAF8")
        set_cell_margins(row.cells[0], top=50, bottom=50, left=80, right=80)
        set_cell_margins(row.cells[1], top=50, bottom=50, left=80, right=80)
        
        p_lbl = row.cells[0].paragraphs[0]
        p_lbl.paragraph_format.space_after = Pt(1)
        r_lbl = p_lbl.add_run(label)
        r_lbl.bold = True
        r_lbl.font.size = Pt(9.5)
        
        p_val = row.cells[1].paragraphs[0]
        p_val.paragraph_format.space_after = Pt(1)
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # About Rivyn
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(8)
    p_sec.paragraph_format.space_after = Pt(2)
    r_sh = p_sec.add_run("About Rivyn")
    r_sh.bold = True
    r_sh.font.size = Pt(12)
    r_sh.font.color.rgb = RGBColor(20, 18, 15)

    p_abt = doc.add_paragraph()
    p_abt.paragraph_format.space_after = Pt(6)
    p_abt.paragraph_format.line_spacing = 1.15
    p_abt.add_run(
        "Rivyn is a next-generation AI log intelligence and observability platform founded out of the prestigious "
        "MHP Hackathon (A Porsche Company) in Germany. We solve the $50B enterprise observability crisis by transforming "
        "gigabytes of messy, unformatted system logs into high-speed columnar Apache Parquet storage and using 3-tier "
        "machine learning to suppress 97%+ of alert noise.\n\n"
        "We are setting up our core engineering division in India and want passionate student developers to join as founding interns."
    )

    # Callout: Important note for freshers
    add_callout_box(
        doc,
        "📢 Important Note for Freshers & Students (Read this first!):",
        "We do NOT expect you to know every single tool or framework listed below! You are a college student / fresher, and we completely respect that. We don't care if you don't know everything on day one. What we care about deeply is your curiosity, design taste, willingness to learn fast, and authentic enthusiasm.\n\n"
        "We are HIGHLY interested in reading your personal motivational letter. Tell us who you are, what you've built, and why you want to work with us.\n\n"
        "⚠️ Please do NOT use ChatGPT or AI to write your letter! We can easily spot AI-generated boilerplate within 5 seconds and we will immediately reject it. We want to hear your real voice, your honest story, and what drives you.",
        bg_hex="FDF8F0",
        border_hex="B45309"
    )

    # About the Role
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(6)
    p_sec.paragraph_format.space_after = Pt(2)
    r_sh = p_sec.add_run("About the Role")
    r_sh.bold = True
    r_sh.font.size = Pt(12)
    r_sh.font.color.rgb = RGBColor(20, 18, 15)

    p_rol = doc.add_paragraph()
    p_rol.paragraph_format.space_after = Pt(6)
    p_rol.paragraph_format.line_spacing = 1.15
    p_rol.add_run(
        "At Rivyn, you will work on real developer-facing tools (think Vercel, Linear, Datadog, Stripe aesthetic). "
        "You will collaborate directly with the founders to build interactive incident boards, dynamic telemetry "
        "timeline charts, and our commercial product showcase website."
    )

    # Responsibilities
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(6)
    p_sec.paragraph_format.space_after = Pt(2)
    r_sh = p_sec.add_run("Selected intern's day-to-day responsibilities include:")
    r_sh.bold = True
    r_sh.font.size = Pt(11)

    resps_p1 = [
        ("Interactive Dashboard Development: ", "Build and polish our live Incident Board, temporal anomaly density graphs, and high-speed log table explorer."),
        ("Product UI & Marketing Website: ", "Enhance our commercial SaaS product portal, interactive ROI calculators, and interactive feature walkthroughs."),
        ("UI/UX Prototyping: ", "Create modern, developer-first designs in Figma and translate them into responsive HTML5, modern CSS3, and JavaScript components."),
        ("Theme & Animation Polish: ", "Build fluid micro-interactions, dark/light theme persistence, and snappy 60 FPS charts."),
        ("Backend Integration: ", "Connect front-end interfaces to our asynchronous Python/FastAPI REST and WebSocket APIs.")
    ]
    for bold_prefix, text in resps_p1:
        p_item = doc.add_paragraph(style='List Bullet')
        p_item.paragraph_format.space_after = Pt(2)
        p_item.paragraph_format.line_spacing = 1.15
        r_b = p_item.add_run(bold_prefix)
        r_b.bold = True
        p_item.add_run(text)

    # Skills We Value
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(6)
    p_sec.paragraph_format.space_after = Pt(2)
    r_sh = p_sec.add_run("Skill(s) We Value (Even basic/intermediate exposure is welcome!):")
    r_sh.bold = True
    r_sh.font.size = Pt(11)

    skills_p1 = [
        ("Core Web: ", "HTML5, Modern CSS3 (Flexbox, CSS Grid, animations, variables), Vanilla JavaScript (ES6+) or TypeScript."),
        ("Design Eye: ", "Clean aesthetic sense, familiarity with Figma or Penpot, and an eye for typography, spacing, and visual hierarchy."),
        ("Data Visualization (Bonus): ", "Exposure to Chart.js, D3.js, Apache ECharts, or Canvas."),
        ("Version Control: ", "Git & GitHub basics.")
    ]
    for bold_prefix, text in skills_p1:
        p_item = doc.add_paragraph(style='List Bullet')
        p_item.paragraph_format.space_after = Pt(2)
        r_b = p_item.add_run(bold_prefix)
        r_b.bold = True
        p_item.add_run(text)

    # Who Can Apply
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(6)
    p_sec.paragraph_format.space_after = Pt(2)
    r_sh = p_sec.add_run("Who Can Apply")
    r_sh.bold = True
    r_sh.font.size = Pt(11)

    apply_p1 = [
        "Currently pursuing a Bachelor’s degree (B.Tech / B.E. / BCA / B.Sc / B.Des) in CS, IT, Design, or related fields.",
        "Have personal projects, hobby websites, or Figma files you created because you genuinely love building.",
        "Can commit 20–30 hours/week with flexibility around your college exam schedule.",
        "Can start immediately or within 1–2 weeks."
    ]
    for item in apply_p1:
        p_item = doc.add_paragraph(style='List Bullet')
        p_item.paragraph_format.space_after = Pt(2)
        p_item.add_run(item)

    # Perks & Benefits
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(6)
    p_sec.paragraph_format.space_after = Pt(2)
    r_sh = p_sec.add_run("Perks & Student Benefits")
    r_sh.bold = True
    r_sh.font.size = Pt(11)

    perks_p1 = [
        ("📜 Internship Certificate & LOR: ", "Formal completion certificate and Letter of Recommendation."),
        ("🚀 Pre-Placement Offer (PPO): ", "Fast-track to a full-time Founding Frontend Engineer role upon graduation."),
        ("🎓 Exam-Friendly & Flexible: ", "Zero penalty for college mid-terms and finals; work hours adapt to your semester schedule."),
        ("📑 College Project / NOC Support: ", "We provide official internship documentation and mentor sign-offs required by your university."),
        ("💡 Founder Mentorship: ", "Direct 1-on-1 mentorship with exposure to European enterprise tech.")
    ]
    for bold_prefix, text in perks_p1:
        p_item = doc.add_paragraph(style='List Bullet')
        p_item.paragraph_format.space_after = Pt(2)
        r_b = p_item.add_run(bold_prefix)
        r_b.bold = True
        p_item.add_run(text)

    # Screening Questions
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(6)
    p_sec.paragraph_format.space_after = Pt(2)
    r_sh = p_sec.add_run("Internshala Screening Questions")
    r_sh.bold = True
    r_sh.font.size = Pt(11)

    questions_p1 = [
        "Share links to your portfolio, GitHub, Figma, or 1–2 websites/dashboards you have built.",
        "Motivational Note (Most Important): In your own genuine words (NO AI/ChatGPT), tell us why you want to join Rivyn and what kind of interfaces you love designing/building. What excites you?",
        "Which college and degree (with graduation year) are you currently pursuing?"
    ]
    for idx, q in enumerate(questions_p1, start=1):
        p_item = doc.add_paragraph(style='List Number')
        p_item.paragraph_format.space_after = Pt(2)
        p_item.add_run(q)

    # PAGE BREAK FOR POSTING 2
    doc.add_page_break()

    # ==================== POSTING 2 ====================
    h1_p2 = doc.add_paragraph()
    h1_p2.paragraph_format.space_before = Pt(4)
    h1_p2.paragraph_format.space_after = Pt(4)
    r_h2 = h1_p2.add_run("Posting 2: Core Developer Intern (Python / AI & Systems)")
    r_h2.font.name = 'Calibri'
    r_h2.font.size = Pt(16)
    r_h2.bold = True
    r_h2.font.color.rgb = RGBColor(15, 118, 110)

    # Meta Table
    tbl_p2 = doc.add_table(rows=6, cols=2)
    tbl_p2.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_p2.autofit = False
    
    meta_p2 = [
        ("Job / Internship Title", "Core Software Developer Intern (Python / Systems & ML)"),
        ("Profile Category", "Software Development / Python Development / Machine Learning"),
        ("Workplace Type", "Work from Home (Remote)"),
        ("Duration", "3 to 6 Months (Flexible around college exams & classes)"),
        ("Stipend", "₹20,000 – ₹35,000 / month (+ Pre-Placement Offer upon graduation)"),
        ("Eligible Degrees & Batches", "B.Tech / B.E. / BCA / B.Sc (CS, IT, AI-DS, ECE) (2025, 2026, 2027 passouts)"),
    ]

    for idx, (label, val) in enumerate(meta_p2):
        row = tbl_p2.rows[idx]
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(4.3)
        
        set_cell_background(row.cells[0], "F2EEE8")
        set_cell_background(row.cells[1], "FAFAF8")
        set_cell_margins(row.cells[0], top=50, bottom=50, left=80, right=80)
        set_cell_margins(row.cells[1], top=50, bottom=50, left=80, right=80)
        
        p_lbl = row.cells[0].paragraphs[0]
        p_lbl.paragraph_format.space_after = Pt(1)
        r_lbl = p_lbl.add_run(label)
        r_lbl.bold = True
        r_lbl.font.size = Pt(9.5)
        
        p_val = row.cells[1].paragraphs[0]
        p_val.paragraph_format.space_after = Pt(1)
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # About Rivyn
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(8)
    p_sec.paragraph_format.space_after = Pt(2)
    r_sh = p_sec.add_run("About Rivyn")
    r_sh.bold = True
    r_sh.font.size = Pt(12)
    r_sh.font.color.rgb = RGBColor(20, 18, 15)

    p_abt2 = doc.add_paragraph()
    p_abt2.paragraph_format.space_after = Pt(6)
    p_abt2.paragraph_format.line_spacing = 1.15
    p_abt2.add_run(
        "Rivyn is an enterprise AI log intelligence platform founded out of the prestigious MHP Hackathon "
        "(A Porsche Company) in Germany. We solve the $50B enterprise observability crisis by transforming gigabytes "
        "of raw system logs into zero-copy columnar Apache Parquet storage, slashing alert fatigue by 97%+ through 3-tier "
        "ML anomaly detection, and providing grounded AI root-cause diagnostics with exact line citations.\n\n"
        "We are hiring student engineers in India who want to work on deep systems programming and applied machine learning."
    )

    # Callout: Important note for freshers
    add_callout_box(
        doc,
        "📢 Important Note for Freshers & Students (Read this first!):",
        "We do NOT expect you to be a master of everything mentioned below! You are still in college, and we don't expect you to have 5 years of industry experience with distributed systems. What matters to us is your algorithmic curiosity, strong fundamentals, and hunger to solve hard problems.\n\n"
        "We are HIGHLY interested in reading your personal motivational letter. Tell us what projects you've tinkered with, what problems fascinate you, and why you want to work on core systems.\n\n"
        "⚠️ Please do NOT use ChatGPT or AI to write your response! We want to read your genuine thoughts. AI-written generic paragraphs will be discarded immediately.",
        bg_hex="FDF8F0",
        border_hex="B45309"
    )

    # About the Role
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(6)
    p_sec.paragraph_format.space_after = Pt(2)
    r_sh = p_sec.add_run("About the Role")
    r_sh.bold = True
    r_sh.font.size = Pt(12)
    r_sh.font.color.rgb = RGBColor(20, 18, 15)

    p_rol2 = doc.add_paragraph()
    p_rol2.paragraph_format.space_after = Pt(6)
    p_rol2.paragraph_format.line_spacing = 1.15
    p_rol2.add_run(
        "Tired of building basic todo apps or generic chatbot wrappers? At Rivyn, you will work on real systems "
        "and data engineering: zero-copy memory mapping, online prefix-tree clustering, SIMD vectorized filters, "
        "and grounded vector embeddings over millions of production log lines."
    )

    # Responsibilities
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(6)
    p_sec.paragraph_format.space_after = Pt(2)
    r_sh = p_sec.add_run("Selected intern's day-to-day responsibilities include:")
    r_sh.bold = True
    r_sh.font.size = Pt(11)

    resps_p2 = [
        ("High-Speed Log Parsing: ", "Optimize our single-pass Drain3 online prefix-tree algorithm to parse unformatted system logs at >15,000 lines/second."),
        ("Columnar Binary Engine: ", "Work with PyArrow and Apache Parquet to implement zero-copy memory mapping, dictionary encoding, and vectorized filter pushdowns."),
        ("ML Anomaly Detection: ", "Implement and evaluate our 3-tier hybrid anomaly detection engine (Statistical template rarity, Scikit-Learn Isolation Forest, and Markov sequence state transition models)."),
        ("Grounded AI Copilot: ", "Enhance our Retrieval-Augmented Generation (RAG) pipeline using dense sentence embeddings (all-MiniLM-L6-v2) and strict citation validation."),
        ("Backend APIs & Testing: ", "Build high-concurrency FastAPI asynchronous REST endpoints and write comprehensive pytest test suites.")
    ]
    for bold_prefix, text in resps_p2:
        p_item = doc.add_paragraph(style='List Bullet')
        p_item.paragraph_format.space_after = Pt(2)
        p_item.paragraph_format.line_spacing = 1.15
        r_b = p_item.add_run(bold_prefix)
        r_b.bold = True
        p_item.add_run(text)

    # Skills We Value
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(6)
    p_sec.paragraph_format.space_after = Pt(2)
    r_sh = p_sec.add_run("Skill(s) We Value (Even basic/intermediate exposure is welcome!):")
    r_sh.bold = True
    r_sh.font.size = Pt(11)

    skills_p2 = [
        ("Core Python: ", "Good grasp of Python basics, data structures, and object-oriented programming."),
        ("Data & Storage (Bonus): ", "Curiosity or exposure to PyArrow, Apache Parquet, DuckDB, or Pandas."),
        ("ML Basics: ", "Fundamental understanding of machine learning concepts (clustering, basic statistics, or embeddings)."),
        ("Backend: ", "Basic experience building an API in FastAPI, Flask, or Django."),
        ("Tools: ", "Basic Git, GitHub, and Linux terminal comfort.")
    ]
    for bold_prefix, text in skills_p2:
        p_item = doc.add_paragraph(style='List Bullet')
        p_item.paragraph_format.space_after = Pt(2)
        r_b = p_item.add_run(bold_prefix)
        r_b.bold = True
        p_item.add_run(text)

    # Who Can Apply
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(6)
    p_sec.paragraph_format.space_after = Pt(2)
    r_sh = p_sec.add_run("Who Can Apply")
    r_sh.bold = True
    r_sh.font.size = Pt(11)

    apply_p2 = [
        "Currently pursuing a Bachelor’s degree (B.Tech / B.E. / BCA / B.Sc) in CS, IT, AI-DS, ECE, or related branches.",
        "Love coding and have personal side-projects, competitive programming, or hackathon code on GitHub.",
        "Can commit 20–30 hours/week with full flexibility for college internals/exams.",
        "Can start immediately or within 1–2 weeks."
    ]
    for item in apply_p2:
        p_item = doc.add_paragraph(style='List Bullet')
        p_item.paragraph_format.space_after = Pt(2)
        p_item.add_run(item)

    # Perks & Benefits
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(6)
    p_sec.paragraph_format.space_after = Pt(2)
    r_sh = p_sec.add_run("Perks & Student Benefits")
    r_sh.bold = True
    r_sh.font.size = Pt(11)

    perks_p2 = [
        ("📜 Internship Certificate & LOR: ", "Formal completion certificate and Letter of Recommendation."),
        ("🚀 Pre-Placement Offer (PPO): ", "Direct conversion into a full-time Founding Systems Engineer upon graduation with competitive package + equity."),
        ("🎓 Exam-Friendly & Flexible: ", "Zero penalty for college mid-terms and finals; work hours adapt to your semester schedule."),
        ("📑 College Project / NOC Support: ", "We support college evaluation requirements, capstone project credits, and internship certificates."),
        ("🧠 Real CS Experience: ", "Direct mentorship from founders on high-performance computing, distributed storage, and applied AI.")
    ]
    for bold_prefix, text in perks_p2:
        p_item = doc.add_paragraph(style='List Bullet')
        p_item.paragraph_format.space_after = Pt(2)
        r_b = p_item.add_run(bold_prefix)
        r_b.bold = True
        p_item.add_run(text)

    # Screening Questions
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(6)
    p_sec.paragraph_format.space_after = Pt(2)
    r_sh = p_sec.add_run("Internshala Screening Questions")
    r_sh.bold = True
    r_sh.font.size = Pt(11)

    questions_p2 = [
        "Share your GitHub profile or links to any code/project you built and are proud of.",
        "Motivational Note (Most Important): In your own genuine words (NO AI/ChatGPT), explain why you want to work on core backend/systems engineering at Rivyn. What was a challenging bug or project you enjoyed figuring out?",
        "Which college and degree/branch are you studying, and what is your expected graduation year?"
    ]
    for idx, q in enumerate(questions_p2, start=1):
        p_item = doc.add_paragraph(style='List Number')
        p_item.paragraph_format.space_after = Pt(2)
        p_item.add_run(q)

    # Pro-tip Box
    add_callout_box(
        doc,
        "💡 Pro-Tip for Filtering on Internshala:",
        "When review time comes, sort applicants by their response to Question 2. Students who write 2–3 authentic, enthusiastic sentences about their own projects will instantly stand out over the hundreds who copy-paste generic AI answers!",
        bg_hex="F0FDF4",
        border_hex="047857"
    )

    # Save document
    out_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs")
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "Rivyn_Internshala_Job_Postings.docx")
    doc.save(out_path)
    print(f"Successfully created Word document: {out_path} ({os.path.getsize(out_path):,} bytes)")

if __name__ == "__main__":
    create_document()
